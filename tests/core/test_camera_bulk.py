"""CCTV 일괄 등록·내보내기 (core/camera_bulk.py).

DB가 필요한 검사는 실제 세션을 쓰고, 엑셀 왕복은 메모리에서 끝낸다.
"""
from __future__ import annotations

import io

import pytest

from tot_dashboard.core import camera_bulk as B
from tot_dashboard.core import cameras as C


@pytest.fixture
def db(db_schema):
    from tot_dashboard.core.db import get_session
    from tot_dashboard.core.models import Camera, CameraDomain, CameraRoi
    s = get_session()
    s.query(CameraRoi).delete()
    s.query(CameraDomain).delete()
    s.query(Camera).delete()
    s.commit()
    yield s
    s.rollback()
    s.close()


def _row(cid="CAM-TEST-1", **over):
    base = {
        "id": cid, "name": "시험지점", "dept": "도로관리과",
        "lat": 35.15, "lng": 129.05,
        "source_type": "hls", "source_url": "https://example.go.kr/a.m3u8",
        "source_path": "", "cctv_name": "시험_1", "is_active": True, "note": "",
        "use_flood": True, "cont_flood": True,
        "use_crowd": False, "cont_crowd": False,
        "use_road": False, "cont_road": False,
    }
    base.update(over)
    return base


# --- ID 생성 ----------------------------------------------------------------
def test_같은_이름과_주소면_항상_같은_ID가_나온다():
    """다시 수집해도 중복 등록되지 않아야 한다."""
    a = B.suggest_id("서면교차로", "https://x/1.m3u8", "BSITS")
    b = B.suggest_id("서면교차로", "https://x/1.m3u8", "BSITS")
    assert a == b


def test_다른_지점은_다른_ID를_받는다():
    a = B.suggest_id("서면교차로", "https://x/1.m3u8")
    b = B.suggest_id("서면교차로", "https://x/2.m3u8")
    assert a != b


def test_생성된_ID가_등록_규칙을_만족한다():
    """한글 이름이라도 ID 규칙(대문자 시작 영숫자)을 지켜야 한다."""
    for name in ("서면교차로", "BEXCO 앞", "제2터널", "123"):
        assert C.ID_RE.match(B.suggest_id(name, "seed")), name


# --- API 결과 변환 ----------------------------------------------------------
def test_API_결과는_탐지서비스가_꺼진_채로_들어온다():
    """전부 켜서 넣으면 CPU가 감당하지 못한다 — 사람이 고르게 둔다."""
    rows = B.from_api_rows([{
        "name": "서면교차로", "lat": 35.15, "lng": 129.05,
        "source_type": "hls", "source_url": "https://x/1.m3u8",
        "cctv_name": "서면교차로"}])
    assert len(rows) == 1
    for d in B.DOMAINS:
        assert rows[0][f"use_{d}"] is False
        assert rows[0][f"cont_{d}"] is False


# --- 검증 -------------------------------------------------------------------
def test_파일_안의_ID_중복을_잡는다(db):
    checked = B.check(db, [_row("CAM-DUPE"), _row("CAM-DUPE")])
    assert checked[0]["_status"] == "new"
    assert checked[1]["_status"] == "error"
    assert any("중복" in e for e in checked[1]["_errors"])


def test_좌표가_국내를_벗어나면_오류다(db):
    checked = B.check(db, [_row(lat=99.0)])
    assert checked[0]["_status"] == "error"


def test_이미_등록된_ID는_오류가_아니라_기존으로_분류한다(db):
    cam, errs = C.create(db, _row("CAM-EXIST"))
    assert not errs
    db.flush()
    checked = B.check(db, [_row("CAM-EXIST")])
    assert checked[0]["_status"] == "dup"


# --- 등록 -------------------------------------------------------------------
def test_기존_ID는_기본적으로_건너뛴다(db):
    """엑셀을 잘못 올려 운영 중인 지점 설정이 지워지는 것이 더 큰 사고다."""
    C.create(db, _row("CAM-KEEP", name="원래이름"))
    db.flush()
    res = B.apply(db, [_row("CAM-KEEP", name="바뀐이름")])
    assert res["skipped"] == 1 and res["created"] == 0
    assert C.get(db, "CAM-KEEP").name == "원래이름"


def test_덮어쓰기를_켜면_갱신한다(db):
    C.create(db, _row("CAM-OVER", name="원래이름"))
    db.flush()
    res = B.apply(db, [_row("CAM-OVER", name="바뀐이름")], update_existing=True)
    assert res["updated"] == 1
    assert C.get(db, "CAM-OVER").name == "바뀐이름"


def test_오류_행은_건너뛰고_나머지는_등록한다(db):
    res = B.apply(db, [_row("CAM-OK-1"), _row("CAM-BAD", lng=999.0)])
    assert res["created"] == 1
    assert res["failed"] == 1
    assert C.get(db, "CAM-OK-1") is not None
    assert C.get(db, "CAM-BAD") is None


def test_엑셀의_사용_상시_열이_도메인_지정에_반영된다(db):
    B.apply(db, [_row("CAM-DOM", use_flood=True, cont_flood=True,
                      use_road=True, cont_road=False)])
    cam = C.get(db, "CAM-DOM")
    dm = C.domain_map(cam)
    assert dm["flood"] == {"enabled": True, "continuous": True, "config": {}}
    assert dm["road"]["enabled"] is True and dm["road"]["continuous"] is False
    assert dm["crowd"]["enabled"] is False


def test_사용하지_않는_도메인은_상시로_남지_않는다(db):
    """쓰지 않는 도메인이 상시로 남아 있으면 파이프라인이 헛돈다."""
    B.apply(db, [_row("CAM-OFF", use_crowd=False, cont_crowd=True)])
    dm = C.domain_map(C.get(db, "CAM-OFF"))
    assert dm["crowd"]["continuous"] is False


# --- 엑셀 -------------------------------------------------------------------
def test_양식에_필수_열이_모두_있다():
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(B.template_excel()))
    header = [c.value for c in wb["CCTV등록양식"][1]]
    assert header == B.HEADERS


def test_내보낸_파일을_그대로_다시_읽을_수_있다(db):
    """받아서 고친 뒤 그대로 올리는 흐름이 성립해야 한다."""
    B.apply(db, [_row("CAM-RT-1", name="왕복시험")])
    db.flush()
    rows, errs = B.parse_excel(B.export_excel(db))
    assert not errs
    got = next(r for r in rows if r["id"] == "CAM-RT-1")
    assert got["name"] == "왕복시험"
    assert got["source_url"] == "https://example.go.kr/a.m3u8"
    assert got["use_flood"] is True and got["cont_flood"] is True


def test_머리글이_다르면_명확히_알려준다():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["엉뚱한", "열"])
    buf = io.BytesIO()
    wb.save(buf)
    rows, errs = B.parse_excel(buf.getvalue())
    assert rows == []
    assert errs and "필수 열" in errs[0]


def test_엑셀이_아닌_파일은_안내를_돌려준다():
    rows, errs = B.parse_excel(b"not an excel file")
    assert rows == []
    assert errs and "열지 못했습니다" in errs[0]


def test_빈_줄은_건너뛴다():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(B.HEADERS)
    ws.append([None] * len(B.HEADERS))
    ws.append(["CAM-A", "가", "", 35.1, 129.0, "hls", "https://x/a.m3u8",
               "", "", "Y", "", "Y", "N", "N", "N", "N", "N"])
    buf = io.BytesIO()
    wb.save(buf)
    rows, errs = B.parse_excel(buf.getvalue())
    assert not errs
    assert [r["id"] for r in rows] == ["CAM-A"]
